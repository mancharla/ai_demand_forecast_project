import os
import pandas as pd

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from sqlalchemy.orm import Session

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart

from app.database import SessionLocal
from app.models.dataset import Dataset
from app.models.forecast import Forecast
from app.models.user import User
from app.models.notification import Notification
from app.utils.security import verify_token


router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)

security = HTTPBearer()

REPORT_FOLDER = "app/reports"

os.makedirs(REPORT_FOLDER, exist_ok=True)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def get_current_user(credentials: HTTPAuthorizationCredentials, db: Session):
    token = credentials.credentials
    email = verify_token(token)

    user = db.query(User).filter(User.email == email).first()

    if not user:
        raise HTTPException(status_code=401, detail="User not found")

    return user


def get_user_forecasts(user: User, db: Session):
    return (
        db.query(Forecast)
        .join(Dataset, Forecast.dataset_id == Dataset.id)
        .filter(Dataset.user_id == user.id)
        .all()
    )


def get_user_dashboard_summary(user: User, db: Session):
    datasets = db.query(Dataset).filter(
        Dataset.user_id == user.id
    ).all()

    total_sales = 0.0
    total_products = set()

    for dataset in datasets:
        if not os.path.exists(dataset.file_path):
            continue

        try:
            if dataset.file_path.endswith(".csv"):
                df = pd.read_csv(dataset.file_path)
            else:
                df = pd.read_excel(dataset.file_path)

            df.columns = df.columns.astype(str).str.strip()

            if "Sales" in df.columns:
                total_sales += float(df["Sales"].sum())

            if "Product" in df.columns:
                for product in df["Product"].unique():
                    total_products.add(str(product))

        except Exception:
            pass

    forecasts = get_user_forecasts(user, db)

    return {
        "user_name": user.name,
        "email": user.email,
        "total_sales": round(total_sales, 2),
        "total_products": len(total_products),
        "forecast_count": len(forecasts),
        "forecasts": forecasts
    }


@router.get("/preview")
def report_preview(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):
    user = get_current_user(credentials, db)

    summary = get_user_dashboard_summary(user, db)

    forecasts = summary["forecasts"]

    forecast_data = [
        {
            "product": forecast.product_name,
            "predicted_sales": float(forecast.predicted_sales)
        }
        for forecast in forecasts
    ]

    return {
        "user_name": summary["user_name"],
        "email": summary["email"],
        "total_sales": summary["total_sales"],
        "total_products": summary["total_products"],
        "forecast_count": summary["forecast_count"],
        "forecast_data": forecast_data
    }


@router.get("/forecast/excel")
def export_forecast_excel(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):
    user = get_current_user(credentials, db)

    forecasts = get_user_forecasts(user, db)

    if not forecasts:
        raise HTTPException(
            status_code=404,
            detail="No forecast data found. Generate forecast first."
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forecast Report"

    sheet.append([
        "Product Name",
        "Predicted Sales"
    ])

    for forecast in forecasts:
        sheet.append([
            str(forecast.product_name),
            float(forecast.predicted_sales)
        ])

    chart = BarChart()
    chart.title = "Product-wise Forecast Sales"
    chart.x_axis.title = "Products"
    chart.y_axis.title = "Predicted Sales"

    data = Reference(
        sheet,
        min_col=2,
        min_row=1,
        max_row=len(forecasts) + 1
    )

    categories = Reference(
        sheet,
        min_col=1,
        min_row=2,
        max_row=len(forecasts) + 1
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    sheet.add_chart(chart, "E2")

    file_path = f"{REPORT_FOLDER}/forecast_report_user_{user.id}.xlsx"

    workbook.save(file_path)

    notification = Notification(
        user_id=user.id,
        title="Excel Report Generated",
        message="Forecast Excel report generated successfully"
    )

    db.add(notification)
    db.commit()

    return FileResponse(
        path=file_path,
        filename="forecast_report.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/dashboard/pdf")
def export_dashboard_pdf(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):
    user = get_current_user(credentials, db)

    summary = get_user_dashboard_summary(user, db)

    forecasts = summary["forecasts"]

    file_path = f"{REPORT_FOLDER}/dashboard_report_user_{user.id}.pdf"

    doc = SimpleDocTemplate(file_path)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(
        Paragraph(
            "AI Demand Forecasting Dashboard Report",
            styles["Title"]
        )
    )

    elements.append(Spacer(1, 20))

    report_data = [
        f"User: {summary['user_name']}",
        f"Email: {summary['email']}",
        f"Total Sales: {summary['total_sales']}",
        f"Total Products: {summary['total_products']}",
        f"Forecast Count: {summary['forecast_count']}"
    ]

    for item in report_data:
        elements.append(
            Paragraph(item, styles["BodyText"])
        )

        elements.append(Spacer(1, 10))

    if forecasts:
        elements.append(Spacer(1, 20))
        elements.append(
            Paragraph(
                "Product-wise Forecast Chart",
                styles["Heading2"]
            )
        )

        drawing = Drawing(450, 250)

        chart = VerticalBarChart()
        chart.x = 50
        chart.y = 40
        chart.height = 160
        chart.width = 350

        values = [
            float(forecast.predicted_sales)
            for forecast in forecasts[:8]
        ]

        labels = [
            str(forecast.product_name)
            for forecast in forecasts[:8]
        ]

        chart.data = [values]
        chart.categoryAxis.categoryNames = labels

        chart.valueAxis.valueMin = 0

        if values:
            chart.valueAxis.valueMax = max(values) + 100
            chart.valueAxis.valueStep = max(50, int(max(values) / 5))

        drawing.add(chart)
        elements.append(drawing)

    doc.build(elements)

    notification = Notification(
        user_id=user.id,
        title="PDF Report Generated",
        message="Dashboard PDF report generated successfully"
    )

    db.add(notification)
    db.commit()

    return FileResponse(
        path=file_path,
        filename="dashboard_report.pdf",
        media_type="application/pdf"
    )