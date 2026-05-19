import os

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from sqlalchemy.orm import Session
from sqlalchemy import or_

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart

from app.database import SessionLocal
from app.models.user import User
from app.models.dataset import Dataset
from app.models.forecast import Forecast
from app.models.notification import Notification
from app.models.forecast_history import ForecastHistory
from app.models.accuracy_metric import AccuracyMetric
from app.utils.security import verify_token


router = APIRouter(
    prefix="/admin",
    tags=["Admin"]
)

security = HTTPBearer()

REPORT_FOLDER = "app/reports"
os.makedirs(REPORT_FOLDER, exist_ok=True)


def get_db():

    db = SessionLocal()

    try:
        yield db

    finally:
        db.close()


def get_admin_user(
    credentials: HTTPAuthorizationCredentials,
    db: Session
):

    token = credentials.credentials
    email = verify_token(token)

    user = db.query(User).filter(
        User.email == email
    ).first()

    if not user:
        raise HTTPException(
            status_code=401,
            detail="User not found"
        )

    if user.role != "admin":
        raise HTTPException(
            status_code=403,
            detail="Admin access required"
        )

    return user


def paginate_query(query, page: int, limit: int):

    total = query.count()

    items = (
        query
        .offset((page - 1) * limit)
        .limit(limit)
        .all()
    )

    return items, total


@router.get("/dashboard")
def admin_dashboard(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    admin = get_admin_user(credentials, db)

    return {
        "admin": {
            "id": admin.id,
            "name": admin.name,
            "email": admin.email,
            "role": admin.role
        },

        "stats": {
            "total_users": db.query(User).count(),
            "total_admins": db.query(User).filter(User.role == "admin").count(),
            "total_normal_users": db.query(User).filter(User.role == "user").count(),
            "total_datasets": db.query(Dataset).count(),
            "total_forecasts": db.query(Forecast).count(),
            "total_notifications": db.query(Notification).count(),
            "total_forecast_history": db.query(ForecastHistory).count(),
            "total_accuracy_metrics": db.query(AccuracyMetric).count()
        },

        "recent_users": [
            {
                "id": user.id,
                "name": user.name,
                "email": user.email,
                "role": user.role
            }
            for user in db.query(User).order_by(User.id.desc()).limit(5).all()
        ],

        "recent_datasets": [
            {
                "id": dataset.id,
                "file_name": dataset.file_name,
                "file_path": dataset.file_path,
                "user_id": dataset.user_id
            }
            for dataset in db.query(Dataset).order_by(Dataset.id.desc()).limit(5).all()
        ],

        "recent_forecasts": [
            {
                "id": forecast.id,
                "product_name": forecast.product_name,
                "predicted_sales": float(forecast.predicted_sales),
                "dataset_id": forecast.dataset_id
            }
            for forecast in db.query(Forecast).order_by(Forecast.id.desc()).limit(5).all()
        ]
    }


@router.get("/users")
def get_all_users(
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
    search: str = "",
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    query = db.query(User)

    if search:
        query = query.filter(
            or_(
                User.name.like(f"%{search}%"),
                User.email.like(f"%{search}%"),
                User.role.like(f"%{search}%")
            )
        )

    query = query.order_by(User.id.desc())

    users, total = paginate_query(query, page, limit)

    return {
        "page": page,
        "limit": limit,
        "total": total,
        "total_pages": (total + limit - 1) // limit,
        "data": [
            {
                "id": user.id,
                "name": user.name,
                "email": user.email,
                "role": user.role
            }
            for user in users
        ]
    }


@router.put("/users/{user_id}/role")
def update_user_role(
    user_id: int,
    role: str,
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    if role not in ["user", "admin"]:
        raise HTTPException(
            status_code=400,
            detail="Role must be user or admin"
        )

    user = db.query(User).filter(
        User.id == user_id
    ).first()

    if not user:
        raise HTTPException(
            status_code=404,
            detail="User not found"
        )

    user.role = role
    db.commit()

    return {
        "message": "User role updated successfully"
    }


@router.delete("/users/{user_id}")
def delete_user(
    user_id: int,
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    admin = get_admin_user(credentials, db)

    if admin.id == user_id:
        raise HTTPException(
            status_code=400,
            detail="Admin cannot delete own account"
        )

    user = db.query(User).filter(
        User.id == user_id
    ).first()

    if not user:
        raise HTTPException(
            status_code=404,
            detail="User not found"
        )

    user_datasets = db.query(Dataset).filter(
        Dataset.user_id == user.id
    ).all()

    for dataset in user_datasets:

        db.query(Forecast).filter(
            Forecast.dataset_id == dataset.id
        ).delete()

        db.query(AccuracyMetric).filter(
            AccuracyMetric.dataset_id == dataset.id
        ).delete()

        db.query(ForecastHistory).filter(
            ForecastHistory.dataset_id == dataset.id
        ).delete()

        if os.path.exists(dataset.file_path):
            os.remove(dataset.file_path)

        db.delete(dataset)

    db.query(Notification).filter(
        Notification.user_id == user.id
    ).delete()

    db.delete(user)
    db.commit()

    return {
        "message": "User deleted successfully"
    }


@router.get("/datasets")
def get_all_datasets(
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
    search: str = "",
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    query = db.query(Dataset)

    if search:
        query = query.filter(
            or_(
                Dataset.file_name.like(f"%{search}%"),
                Dataset.file_path.like(f"%{search}%")
            )
        )

    query = query.order_by(Dataset.id.desc())

    datasets, total = paginate_query(query, page, limit)

    return {
        "page": page,
        "limit": limit,
        "total": total,
        "total_pages": (total + limit - 1) // limit,
        "data": [
            {
                "id": dataset.id,
                "file_name": dataset.file_name,
                "file_path": dataset.file_path,
                "user_id": dataset.user_id
            }
            for dataset in datasets
        ]
    }


@router.delete("/datasets/{dataset_id}")
def admin_delete_dataset(
    dataset_id: int,
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    dataset = db.query(Dataset).filter(
        Dataset.id == dataset_id
    ).first()

    if not dataset:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found"
        )

    db.query(Forecast).filter(
        Forecast.dataset_id == dataset.id
    ).delete()

    db.query(AccuracyMetric).filter(
        AccuracyMetric.dataset_id == dataset.id
    ).delete()

    db.query(ForecastHistory).filter(
        ForecastHistory.dataset_id == dataset.id
    ).delete()

    if os.path.exists(dataset.file_path):
        os.remove(dataset.file_path)

    db.delete(dataset)
    db.commit()

    return {
        "message": "Dataset deleted successfully"
    }


@router.get("/forecasts")
def get_all_forecasts(
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
    search: str = "",
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    query = db.query(Forecast)

    if search:
        query = query.filter(
            Forecast.product_name.like(f"%{search}%")
        )

    query = query.order_by(Forecast.id.desc())

    forecasts, total = paginate_query(query, page, limit)

    return {
        "page": page,
        "limit": limit,
        "total": total,
        "total_pages": (total + limit - 1) // limit,
        "data": [
            {
                "id": forecast.id,
                "product_name": forecast.product_name,
                "predicted_sales": float(forecast.predicted_sales),
                "dataset_id": forecast.dataset_id
            }
            for forecast in forecasts
        ]
    }


@router.get("/forecast-history")
def get_all_forecast_history(
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
    search: str = "",
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    query = db.query(ForecastHistory)

    if search:
        query = query.filter(
            or_(
                ForecastHistory.model_name.like(f"%{search}%"),
                ForecastHistory.top_demand_product.like(f"%{search}%")
            )
        )

    query = query.order_by(ForecastHistory.id.desc())

    history, total = paginate_query(query, page, limit)

    return {
        "page": page,
        "limit": limit,
        "total": total,
        "total_pages": (total + limit - 1) // limit,
        "data": [
            {
                "id": item.id,
                "user_id": item.user_id,
                "dataset_id": item.dataset_id,
                "model_name": item.model_name,
                "forecast_days": item.forecast_days,
                "top_demand_product": item.top_demand_product,
                "created_at": str(item.created_at)
            }
            for item in history
        ]
    }


@router.get("/accuracy-metrics")
def get_all_accuracy_metrics(
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
    search: str = "",
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    query = db.query(AccuracyMetric)

    if search:
        query = query.filter(
            AccuracyMetric.model_name.like(f"%{search}%")
        )

    query = query.order_by(AccuracyMetric.id.desc())

    metrics, total = paginate_query(query, page, limit)

    return {
        "page": page,
        "limit": limit,
        "total": total,
        "total_pages": (total + limit - 1) // limit,
        "data": [
            {
                "id": item.id,
                "user_id": item.user_id,
                "dataset_id": item.dataset_id,
                "model_name": item.model_name,
                "mae": item.mae,
                "rmse": item.rmse,
                "mape": item.mape,
                "created_at": str(item.created_at)
            }
            for item in metrics
        ]
    }


@router.get("/reports")
def get_uploaded_reports(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    return [
        {
            "name": "All Users Forecast Excel Report",
            "type": "Excel",
            "endpoint": "/admin/reports/forecast/excel"
        },
        {
            "name": "Admin Dashboard PDF Report",
            "type": "PDF",
            "endpoint": "/admin/reports/dashboard/pdf"
        }
    ]


@router.get("/reports/forecast/excel")
def admin_export_all_forecast_excel(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    forecasts = db.query(Forecast).order_by(
        Forecast.id.desc()
    ).all()

    if not forecasts:
        raise HTTPException(
            status_code=404,
            detail="No forecast data found"
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All Forecasts"

    sheet.append([
        "Forecast ID",
        "Product Name",
        "Predicted Sales",
        "Dataset ID"
    ])

    for forecast in forecasts:
        sheet.append([
            int(forecast.id),
            str(forecast.product_name),
            float(forecast.predicted_sales),
            int(forecast.dataset_id)
        ])

    chart = BarChart()
    chart.title = "All Users Product Forecasts"
    chart.x_axis.title = "Products"
    chart.y_axis.title = "Predicted Sales"

    data = Reference(
        sheet,
        min_col=3,
        min_row=1,
        max_row=len(forecasts) + 1
    )

    categories = Reference(
        sheet,
        min_col=2,
        min_row=2,
        max_row=len(forecasts) + 1
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    sheet.add_chart(chart, "F2")

    file_path = f"{REPORT_FOLDER}/admin_all_forecast_report.xlsx"

    workbook.save(file_path)

    return FileResponse(
        path=file_path,
        filename="admin_all_forecast_report.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/reports/dashboard/pdf")
def admin_export_dashboard_pdf(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):

    get_admin_user(credentials, db)

    total_users = db.query(User).count()
    total_admins = db.query(User).filter(User.role == "admin").count()
    total_normal_users = db.query(User).filter(User.role == "user").count()
    total_datasets = db.query(Dataset).count()
    total_forecasts = db.query(Forecast).count()
    total_notifications = db.query(Notification).count()
    total_history = db.query(ForecastHistory).count()
    total_metrics = db.query(AccuracyMetric).count()

    forecasts = db.query(Forecast).order_by(
        Forecast.id.desc()
    ).limit(10).all()

    file_path = f"{REPORT_FOLDER}/admin_dashboard_report.pdf"

    doc = SimpleDocTemplate(file_path)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(
        Paragraph(
            "Admin System Analytics Report",
            styles["Title"]
        )
    )

    elements.append(Spacer(1, 20))

    report_data = [
        f"Total Users: {total_users}",
        f"Total Admins: {total_admins}",
        f"Total Normal Users: {total_normal_users}",
        f"Total Datasets: {total_datasets}",
        f"Total Forecasts: {total_forecasts}",
        f"Total Notifications: {total_notifications}",
        f"Forecast History Records: {total_history}",
        f"Accuracy Metrics Records: {total_metrics}"
    ]

    for item in report_data:
        elements.append(
            Paragraph(item, styles["BodyText"])
        )

        elements.append(Spacer(1, 10))

    if forecasts:
        elements.append(Spacer(1, 20))

        elements.append(
            Paragraph(
                "Recent Forecast Analytics Chart",
                styles["Heading2"]
            )
        )

        drawing = Drawing(450, 250)

        chart = VerticalBarChart()
        chart.x = 50
        chart.y = 40
        chart.height = 160
        chart.width = 350

        values = [
            float(forecast.predicted_sales)
            for forecast in forecasts
        ]

        labels = [
            str(forecast.product_name)
            for forecast in forecasts
        ]

        chart.data = [values]
        chart.categoryAxis.categoryNames = labels
        chart.valueAxis.valueMin = 0

        if values:
            chart.valueAxis.valueMax = max(values) + 100
            chart.valueAxis.valueStep = max(
                50,
                int(max(values) / 5)
            )

        drawing.add(chart)
        elements.append(drawing)

    doc.build(elements)

    return FileResponse(
        path=file_path,
        filename="admin_dashboard_report.pdf",
        media_type="application/pdf"
    )